from openpyxl import load_workbook
wb = load_workbook("sample.xlsx")
ws = wb.active

from openpyxl.chart import BarChart, Reference
bar_value = Reference(ws, min_row = 2, max_row = 11, min_col = 2, max_col = 3) # 영어, 수학 성적을 chart로 그리기 위해 영역 지정
bar_chart = BarChart() # 차트 종류 설정 (Bar, Line, Pie, ...)
bar_chart.add_data(bar_value)
ws.add_chart(bar_chart, "E1") # E1에 차트 넣기(위치 선정)

# B1:C11까지의 데이터 + 제목 포함
line_value = Reference(ws, min_row = 1, max_row = 11, min_col = 2, max_col = 3)
line_chart = LineChart()
line_chart.add_data(line_value, titles_from_data = True) # 계열 -> 영어, 수학 (제목에서 가져옴)
line_chart.title = "성적표" # 제목
line_chart.style = 20 # 미리 적용된 스타일을 적용, 사용자 지정 가능
line_chart.y_axis.title = "점수" # y축의 제목
line_chart.x_axis.title = "번호" # x축의 제목

wb.save("sample_chart.xlsx")